#!/usr/bin/env python3
"""Build the review spreadsheet from whatever the site currently says.

    python3 scripts/build_review_sheet.py

Writes «Άγιος_Νεκτάριος_κείμενα.xlsx» next to the repo root. Upload it to Google
Drive and open it — Drive converts it to a Google Sheet, and anyone with the
link can then propose wording in the yellow columns and comment on any row.

The page texts come from content/legacy/*.html; the home page is read live from
the deployed site, because the admin can edit it and the store is the truth.
Re-run this whenever the texts have moved on and hand out a fresh copy.
"""
import glob
import html
import json
import os
import re
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SITE = "https://agios-nektarios.netlify.app"


def page_titles():
    """The human page names, read straight out of the app so they never drift."""
    src = open(os.path.join(ROOT, "src", "lib", "content", "index.ts"), encoding="utf-8").read()
    body = re.search(r"export const PAGE_TITLES[^{]*\{(.*?)\n\};", src, re.S).group(1)
    return dict(re.findall(r'"?([A-Za-z_\-]+)"?\s*:\s*"([^"]+)"', body))


TITLES = page_titles()
"""Pull every editable sentence out of the site into reviewable rows."""


DIR = os.path.join(ROOT, "content", "legacy")
SITE = "https://agios-nektarios.netlify.app"


BLOCK = re.compile(
    r"<(h1|h2|h3|p|li|figcaption|blockquote)\b([^>]*)>(.*?)</\1>", re.S | re.I)
KIND = {"h1":"Τίτλος", "h2":"Επικεφαλίδα", "h3":"Υπο-επικεφαλίδα", "p":"Παράγραφος",
        "li":"Στοιχείο λίστας", "figcaption":"Λεζάντα", "blockquote":"Παράθεμα"}

def plain(fragment):
    t = re.sub(r"<span class=\"lp-credit\">.*?</span>", "", fragment, flags=re.S)
    t = re.sub(r"<[^>]+>", "", t)
    return re.sub(r"\s+", " ", html.unescape(t)).strip()

def rows():
    out = []
    for path in sorted(glob.glob(f"{DIR}/*.html")):
        slug = os.path.basename(path)[:-5]
        page = TITLES.get(slug, slug)
        body = open(path, encoding="utf-8").read()
        # Archive grids are link labels ("Πρακτικό 7 · PDF"), not prose to review.
        body = re.sub(r'<ul class="lp-grid">.*?</ul>', "", body, flags=re.S)
        n = 0
        for m in BLOCK.finditer(body):
            tag, attrs, inner = m.group(1).lower(), m.group(2), m.group(3)
            text = plain(inner)
            if not text:
                continue
            kind = KIND.get(tag, tag)
            if "lp-kicker" in attrs:  kind = "Ετικέτα"
            if "lp-lede"   in attrs:  kind = "Εισαγωγή"
            if "lp-note"   in attrs:  kind = "Σημείωση πηγών"
            if "lp-credit" in attrs:  continue
            n += 1
            out.append({
                "id": f"{slug}-{n:02d}", "slug": slug, "page": page, "kind": kind,
                "text": text, "html": re.sub(r"\s+", " ", inner).strip(),
                "url": f"{SITE}/{'' if slug=='index' else slug}",
            })
    return out



# ── the home page, read from the live site ──────────────────────────────────
HOME_LABELS = {"hero":"Επικεφαλίδα (hero)","stats":"Στατιστικά","essence":"Στροφή",
  "settlement":"Ο οικισμός μας","folklore":"Κατάλογος θεμάτων","event":"Ο κύκλος της χρονιάς",
  "board":"Διοικητικό Συμβούλιο","documents":"Πρακτικά","news":"Νέα","contact":"Επικοινωνία"}
HOME_FIELDS = {"metaLeft":"Μικρό κείμενο αριστερά","metaCenter":"Μικρό κείμενο κέντρο",
  "metaRight":"Μικρό κείμενο δεξιά","titleTop":"Τίτλος, πάνω γραμμή",
  "titleBottom":"Τίτλος, κάτω γραμμή","intro":"Εισαγωγή","ticker":"Κυλιόμενο κείμενο",
  "label":"Ετικέτα","title":"Τίτλος","text":"Κείμενο","paragraphs":"Παράγραφος",
  "lines":"Στίχος","attribution":"Απόδοση","moreLabel":"Κείμενο συνδέσμου",
  "linkLabel":"Κείμενο συνδέσμου","items":"Στοιχείο","members":"Μέλος",
  "links":"Σύνδεσμος","email":"Email","value":"Τιμή","name":"Όνομα","role":"Ρόλος",
  "sub":"Υπότιτλος","num":"Αριθμός"}
HOME_SKIP = {"href","external","mapEmbedUrl","mapLinkUrl","moreHref","linkHref","imageUrl"}


def live_home():
    """Lift the published home content out of the page's React payload."""
    page = urllib.request.urlopen(urllib.request.Request(
        SITE + "/", headers={"User-Agent": "review-sheet/1.0"}), timeout=60).read().decode()
    chunks = re.findall(r'self\.__next_f\.push\(\[1,("(?:[^"\\]|\\.)*")\]\)', page)
    flat = "".join(json.loads(c) for c in chunks)
    dec, full, hero = json.JSONDecoder(), None, None
    for m in re.finditer(r'\{"hero":', flat):
        try:
            cand, _ = dec.raw_decode(flat[m.start():])
        except Exception:
            continue
        if not isinstance(cand, dict):
            continue
        if "essence" in cand and "folklore" in cand:
            full = full or cand
        elif isinstance(cand.get("hero"), dict) and "titleTop" in cand["hero"]:
            hero = hero or cand["hero"]
    if not (full and hero):
        raise SystemExit("Δεν βρέθηκε το περιεχόμενο της αρχικής στο live σάιτ.")
    full["hero"] = hero
    return full


def home_rows():
    out = []

    def add(section, field, value, path):
        if isinstance(value, str) and value.strip():
            out.append({"id": path, "section": HOME_LABELS.get(section, section),
                        "field": field, "text": value})

    for sec, val in live_home().items():
        if not isinstance(val, dict):
            continue
        for k, v in val.items():
            if k in HOME_SKIP:
                continue
            if isinstance(v, str):
                add(sec, HOME_FIELDS.get(k, k), v, f"home.{sec}.{k}")
            elif isinstance(v, list):
                for i, item in enumerate(v, 1):
                    if isinstance(item, str):
                        add(sec, f"{HOME_FIELDS.get(k, k)} {i}", item, f"home.{sec}.{k}[{i}]")
                    elif isinstance(item, dict):
                        for kk, vv in item.items():
                            if kk in HOME_SKIP or not isinstance(vv, str):
                                continue
                            add(sec, f"{HOME_FIELDS.get(k, k)} {i} — {HOME_FIELDS.get(kk, kk)}",
                                vv, f"home.{sec}.{k}[{i}].{kk}")
    return out



OUT = os.path.join(ROOT, "Άγιος_Νεκτάριος_κείμενα.xlsx")
BODY  = Font(name="Arial", size=10)
HEAD  = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Arial", size=16, bold=True, color="1B1B1B")
MUTED = Font(name="Arial", size=10, color="6B6B6B")
INK   = PatternFill("solid", fgColor="2A3038")           # header band
FILLIN= PatternFill("solid", fgColor="FFF6DC")           # the columns reviewers write in
ZEBRA = PatternFill("solid", fgColor="F5F3EE")
EDGE  = Border(bottom=Side(style="thin", color="D8D4CC"))
WRAP  = Alignment(vertical="top", wrap_text=True)
TOP   = Alignment(vertical="top")

STATUS = '"—,Πρόταση,Συμφωνώ,Διαφωνώ,Έγινε"'

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HEAD, INK
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    # A cell *reference*, not ws.cell() — that would materialise an empty row 2
    # and push every data row down one, silently misaligning the formulas.
    ws.freeze_panes = f"A{row + 1}"

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def add_table(ws, headers, rows, fill_from, link_col=None):
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for i, r in enumerate(rows, start=2):
        ws.append(r)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=c)
            cell.font, cell.border = BODY, EDGE
            cell.alignment = WRAP if c >= 3 else TOP
            if c >= fill_from:
                cell.fill = FILLIN
            elif i % 2 == 0:
                cell.fill = ZEBRA
        if link_col:
            cell = ws.cell(row=i, column=link_col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(name="Arial", size=9, color="1155CC", underline="single")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

wb = Workbook()

# ── 1. Οδηγίες ──────────────────────────────────────────────────────────────
g = wb.active; g.title = "Οδηγίες"
g.sheet_view.showGridLines = False
widths(g, {"A": 3, "B": 96})
lines = [
    ("Άγιος Νεκτάριος — τα κείμενα του ιστότοπου", TITLE),
    ("", None),
    ("Κάθε γραμμή είναι ένα κομμάτι κειμένου όπως εμφανίζεται σήμερα στο agios-nektarios.netlify.app.", BODY),
    ("", None),
    ("Πώς προτείνετε αλλαγή", Font(name="Arial", size=11, bold=True)),
    ("1.  Βρείτε τη γραμμή. Τα φύλλα «Αρχική» και «Σελίδες» έχουν φίλτρα στις επικεφαλίδες.", BODY),
    ("2.  Γράψτε τη διατύπωσή σας στη στήλη «Πρόταση αλλαγής» — τα κίτρινα κελιά είναι δικά σας.", BODY),
    ("3.  Εξηγήστε γιατί στο «Σχόλιο» και βάλτε το όνομά σας στο «Ποιος».", BODY),
    ("4.  Αφήστε τη στήλη «Τρέχον κείμενο» ανέπαφη — είναι το σημείο αναφοράς.", BODY),
    ("", None),
    ("Η στήλη «Κατάσταση» έχει έτοιμες επιλογές: Πρόταση, Συμφωνώ, Διαφωνώ, Έγινε.", BODY),
    ("Έτσι βλέπει κανείς με μια ματιά τι εκκρεμεί.", BODY),
    ("", None),
    ("Ο «Κωδικός» είναι η ταυτότητα κάθε κομματιού. Μην τον αλλάξετε: από αυτόν βρίσκουμε", BODY),
    ("πού ακριβώς πάει η αλλαγή σας μέσα στη σελίδα.", BODY),
    ("", None),
    ("Αν το κείμενο περιέχει σύνδεσμο ή έντονα γράμματα, η στήλη «Σημ.» το λέει. Γράψτε τότε", BODY),
    ("απλά τα λόγια σας — τη μορφοποίηση την ξαναβάζουμε εμείς. Το φύλλο «Τεχνικά» κρατά", BODY),
    ("τον αρχικό κώδικα για όποιον εφαρμόσει τις αλλαγές.", BODY),
]
for text, font in lines:
    g.append(["", text])
    if font: g.cell(row=g.max_row, column=2).font = font

blocks = rows()
home   = home_rows()

def flags(raw):
    marks = []
    if "<a " in raw: marks.append("σύνδεσμος")
    if "<strong" in raw: marks.append("έντονα")
    if "<em" in raw: marks.append("πλάγια")
    return ", ".join(marks)

# ── 2. Αρχική ───────────────────────────────────────────────────────────────
h = wb.create_sheet("Αρχική")
h.sheet_view.showGridLines = False
add_table(h,
    ["Κωδικός","Ενότητα","Πεδίο","Τρέχον κείμενο","Πρόταση αλλαγής","Σχόλιο","Ποιος","Κατάσταση","Λέξεις"],
    [[r["id"], r["section"], r["field"], r["text"], "", "", "", "—", ""] for r in home],
    fill_from=5)
widths(h, {"A":26,"B":22,"C":26,"D":62,"E":62,"F":36,"G":14,"H":13,"I":8})

# ── 3. Σελίδες ──────────────────────────────────────────────────────────────
s = wb.create_sheet("Σελίδες")
s.sheet_view.showGridLines = False
add_table(s,
    ["Κωδικός","Σελίδα","Τύπος","Τρέχον κείμενο","Πρόταση αλλαγής","Σχόλιο","Ποιος","Κατάσταση","Λέξεις","Σημ.","Σύνδεσμος"],
    [[b["id"], b["page"], b["kind"], b["text"], "", "", "", "—", "", flags(b["html"]), b["url"]]
     for b in blocks],
    fill_from=5, link_col=11)
widths(s, {"A":24,"B":24,"C":17,"D":66,"E":66,"F":36,"G":14,"H":13,"I":8,"J":18,"K":30})

# Word count of whatever currently stands — the proposal if there is one, else the original.
for ws, col, n in ((h, "I", len(home)), (s, "I", len(blocks))):
    for r in range(2, n + 2):
        ws[f"{col}{r}"] = (f'=IF(LEN(TRIM(IF(E{r}<>"",E{r},D{r})))=0,0,'
                           f'LEN(TRIM(IF(E{r}<>"",E{r},D{r})))'
                           f'-LEN(SUBSTITUTE(TRIM(IF(E{r}<>"",E{r},D{r}))," ",""))+1)')

for ws, n in ((h, len(home)), (s, len(blocks))):
    dv = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{n + 1}")

# ── 4. Τεχνικά ──────────────────────────────────────────────────────────────
t = wb.create_sheet("Τεχνικά")
t.sheet_view.showGridLines = False
add_table(t, ["Κωδικός","Αρχείο","Αρχικός κώδικας (HTML)"],
          [[b["id"], f"content/legacy/{b['slug']}.html", b["html"]] for b in blocks],
          fill_from=99)
widths(t, {"A":24,"B":34,"C":120})
t["E1"] = "Μόνο για όποιον εφαρμόσει τις αλλαγές. Δεν χρειάζεται να το πειράξει κανείς άλλος."
t["E1"].font = MUTED

# ── σύνοψη στις Οδηγίες, ζωντανή ────────────────────────────────────────────
g.append([""]); g.append(["", "Πού βρισκόμαστε"])
g.cell(row=g.max_row, column=2).font = Font(name="Arial", size=11, bold=True)
base = g.max_row
for label, formula in [
    ("Κομμάτια κειμένου συνολικά", f'=COUNTA(Αρχική!A2:A{len(home)+1})+COUNTA(Σελίδες!A2:A{len(blocks)+1})'),
    ("Με πρόταση αλλαγής",          f'=COUNTIF(Αρχική!H2:H{len(home)+1},"Πρόταση")+COUNTIF(Σελίδες!H2:H{len(blocks)+1},"Πρόταση")'),
    ("Συμφωνήθηκαν",                f'=COUNTIF(Αρχική!H2:H{len(home)+1},"Συμφωνώ")+COUNTIF(Σελίδες!H2:H{len(blocks)+1},"Συμφωνώ")'),
    ("Εφαρμόστηκαν",                f'=COUNTIF(Αρχική!H2:H{len(home)+1},"Έγινε")+COUNTIF(Σελίδες!H2:H{len(blocks)+1},"Έγινε")'),
    ("Λέξεις σε όλο τον ιστότοπο",  f'=SUM(Αρχική!I2:I{len(home)+1})+SUM(Σελίδες!I2:I{len(blocks)+1})'),
]:
    g.append(["", label, formula])
    g.cell(row=g.max_row, column=2).font = BODY
    c = g.cell(row=g.max_row, column=3); c.font = Font(name="Arial", size=10, bold=True)
g.column_dimensions["C"].width = 12

wb.save(OUT)
print("γράφτηκε:", OUT)
